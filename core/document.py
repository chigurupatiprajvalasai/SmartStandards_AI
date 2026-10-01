def extract_uploaded_text(uploaded_file):
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:
            return f"PDF extraction error: {e}"
    if name.endswith(".docx"):
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            return f"DOCX extraction error: {e}"
    return ""
